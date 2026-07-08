from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.utils.cell import range_boundaries


def _sync_table_headers(worksheet) -> None:
    for table in worksheet.tables.values():
        min_col, min_row, max_col, _max_row = range_boundaries(table.ref)
        header_values: list[str] = []
        for col_idx in range(min_col, max_col + 1):
            value = worksheet.cell(row=min_row, column=col_idx).value
            header_values.append(str(value).strip() if value is not None else "")

        for index, column in enumerate(table.tableColumns):
            header_value = header_values[index] if index < len(header_values) else ""
            if header_value:
                column.name = header_value


def export_excel(original_bytes: bytes, parsed_segments: list[dict], request_segments: list[dict]) -> bytes:
    workbook = openpyxl.load_workbook(BytesIO(original_bytes))
    renamed_sheets: dict[str, str] = {}
    for parsed, request in zip(parsed_segments, request_segments):
        translation = (request.get("translation") or request.get("draftTranslation") or request.get("draft_translation") or "").strip()
        if not translation:
            continue
        original_sheet_name = renamed_sheets.get(parsed["sheet"], parsed["sheet"])
        sheet = workbook[original_sheet_name]
        if parsed.get("segment_type") == "sheet_title":
            sheet.title = translation
            renamed_sheets[parsed["sheet"]] = translation
            continue
        sheet[parsed["cell"]] = translation

    for sheet in workbook.worksheets:
        _sync_table_headers(sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
